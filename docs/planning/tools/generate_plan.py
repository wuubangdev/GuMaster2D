"""Run with the bundled Windows Python (openpyxl + python-docx).
Rebuilds planning artifacts from planning_data.py; never edits the user's template.
"""
from pathlib import Path
from datetime import timedelta, date
from collections import defaultdict
from functools import lru_cache
import json, re, math, zipfile, unicodedata, xml.etree.ElementTree as ET
import planning_data as d
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule, CellIsRule
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as col
from openpyxl.utils.datetime import to_excel
from openpyxl.chart import BarChart, Reference

OUT = Path(__file__).resolve().parents[1]
DOCS = OUT.parent
NAVY='142D4E'; TEAL='007E87'; GOLD='E8B44C'; INK='24354B'; GRAY='61738B'
LIGHT='EAF3F5'; BLUE='1D66AC'; RED='FCE4E1'; GREEN='E2F1E8'
VN_SYLLABLES = '''a ai alpha an anh áp ảnh ban bản bán bám bảng bậc bắt bày bền bên biết biên biến bộ bổ bốn bước bởi bớt bình buộc bằng bọc bóng bảo cao cài cái cả cản cảnh cách cấp cần cầu cây chặng chậm chặn chạy chéo chế chết chính chi chịu chia chiến chỉ chữ chưa chuẩn chuỗi chốt chọn cho chơi chờ chuyển chứa chứng chủ chung chương có cũ cũng củ cố cùng còn công cộng cổ cổng cắt cấu cầm cấm câu cân cam càng cuối của cử cửa cứ cụ cực cửa cơ cốt danh dao dài dạng đạt đầu đan đa đã đá đà đẩy đầy đích điểm đi định đọc đồ đội đó đói đối đổi đồng đóng động đột đúng được đừng đủ đưa đường đuổi dùng dụng dự dừng dễ em ép ghi ghép giữ giới giây giờ giảm giá giả giao giải gồm gói gần gắn góc gỡ gốc góp gọn hai hành hàng hạn hãy học hỏa họa hóa hồi hỏi hơn hợp hỏng hoạt hiệu hiểu hiển hiện hệt hết hủy hướng hữu hình hơi hằng hạ hỗ hố hồn hứa hợp hội hoang họ hôm hằng ích ít import kèm kẹt kéo khi khóa khó khỏi khớp không khung khu khuyết khai khác khả kết kiểm kiểu ký kỳ kịch kích kỹ khối khỏe khách làm lại lần lấy lớn lợi lỗi luật luyện lương lượng lặp lên lệ lo lộ lôi loại liên lịch linh lưu lúc luồng lý lượt liệu mời mỗi mới một mất màn máy mạng mặc mạnh mẫu mềm miễn minh mô mốc mở mã mức mục mua màu mượn múi mặt mong này năm nằm nặng nào nền nếu nên người ngày ngại ngẫu ngân nghe nghỉ ngỏ ngoài nguồn ngừng ngay nháp nhân nhập nhận nhỏ nhiều nhờ nhất nhị nhau nhanh như nhưng những nhấn nhiệm nhóm nhìn nhạy nhầm nhịp nhớ nâng nạp nam năm nội nở nơi nổ nữa ổ ô ở phép phải phần phạt phân phẩm phát phòng phôi phổ phút phụ phù phụt phục phức quyền quan quản quái quá qua quay quang quả quỹ quy quên ra rà rào rank rằng rõ rút riêng rời rơi rồi rỗng rộng rừng rớt rủi ro reset rác số sống sớm so sói sử sức sẵn sửa sai sau sâu sẽ sơ sự suất xuất suốt suy sát sách sân săn sản sáng sao sạch sơn sinh sát sự sạp tắt tạo tải tạm tại tay tập tách tài tên thêm thế thể thiết thiếu thiệu thiện tiến tiền tiện tiêu tiếp tích tình tính tự từ tư tử tương tối tối ưu từng tất tăng tầng thật thấy thăng thấp thất thẻ thoại thời thôi thừa thực thử thức thủ thuộc thông thờ thường thước thứ thỏa trước trên trễ trò trả trái trại trùng trường truy truyền trong trục trọn trạng trở tra trợ trọng trữ thu thúc thua thụ thân thẳng thành tháng thanh thấp thị thính thọ thảo tông tồn tô toàn tuần tơ túi tốt tới tổ tuổi tám tấn tầm u ủy và vài vào vẫn vắng vận vẽ về viên việc với vừa vùng vật vững vốn vòng vỏ vui vượt vô văn võ vị vĩnh xóa xếp xem xác xong xử xung xanh xấu xâm xài y yếu ý yêu một lẻ hệ hạt bù dự tái xạ sương ruột vuốt nhãn vách quạt xương thằn lằn giảm độc xám cốt tằm đá trắng vệ bạch ngọc kim chung du cương hạc băng hóa thạch hồi xuân địa thính'''.split()
EN_WORDS = '''account active actor action admin after all alive ally already animation api append arena armor asset assets atlas atomic attack audit auth authority auto avatar available awareness back backend backup bag balance ban base batch before bind block blocker boss bound browser budget buff build bundle burst buy cache callback callbacks camera cancel cap capacity capturable capture case cast catalog channel character chat check checked checkpoint checks checksums claim claimed class cleanup click client closed code codex collection collision combat command commit completion component config conflict connection consume content control cooldown core correction cost counter cpu craft crash create credits csv cumulative currency damage dashboard data database date daily day dead death debug decline default definition delay delete delta demo denied deployment derived destination detail dev dialog direction disconnect dispose doc docs drop dungeon duplicate durable effect elite empty end endpoint engine enter entity entry equip equipped error escaped event evidence exactly exception expiry expired fail failed false farm feature feed fence fencing first fixed flag float focus food foreground foreign fps frame freeze from frontend full function game gate gm go grant grid group guard gu gui guide handoff happy hardware hash heal health hit hook host hosting hot hotkey hp html http https hunger icon id idempotency immutable inbox income input instance intent interactive interrupt inventory invite item jitter job join json key keyboard kick kill known language latency layer layout leader lease leave ledger level limit link list load loader loadout local lock locked log login logout look loop loss low mana map marker mask master match material max memory membership message metric metrics migrate migration min mini minimap mirror missing mobile mode mod moderation modify money monitor mouse move movement music mute native nectar network new no normal not npc null number object offline old one online only opt optin orphan outcome overflow oversize owner ownership page party pass passed passive password patch path pattern payload penalty pending percent performance personal phase pipeline pity pixel player point policy pool portal position post power prefab preferences preview primitive private process processed profile progress projectile prototype publish query quota radius random range rank rate raw read ready real realtime reason receipt recipient reconnect recovery recolor reference refresh regen reject rejection release remap remove render repair replay report request reserve reset resize resolver resource respawn result retry return revoke reward right risk role roll rollback room roster row run running safe same sanitize save scale scene schema seal search second seed self seq sequence server session settings shard shield shop side signature simulation single skill slot slow snapshot socket source spam spawn speed stack stale starter start state static stats status step stone stop stun success support swap sync tag target task tcp terminal test tester text tick tile tileset time timed timer timeout token tools trace tracker trade traffic transaction transfer true tutorial type typecheck ui unverified update use user uuid validate validation value variant version view vision vote vs wait walk warning warp way websocket week weekly win windup won work world write writer wrong xp xss zero zone cleanup exactlyonce online offline fulltime longtime lifetime oneinoneout sourceoftruth lowgraphics frame time frameTime eventId before after sameID maxMana PvP PvE RPG MMORPG TypeScript JavaScript Phaser Colyseus PostgreSQL Redis Node WBS FSD GDD MVP LAN LANalpha RTT CCU RTO RPO CI RC QA QWER GM_READ GM_MOD GM_OWNER NoData NoGo Should Must Won Idle Aggro Chase Attack Leash Dead Main Side Prototype Guard While Until StartTime EndTime ConfigVersion RequestId RoomId CharacterId GuId ItemId ActorId Manifest'''.split()
LEX=set(x.lower() for x in VN_SYLLABLES+EN_WORDS)
LEX.update(['callback','callbacks','safespawn','serverTime','eventId','requestId','characterId','configVersion','maxHP','maxMana','targetMode','castTime','endTime','expectedVersion','inventoryVersion','mapId','roomId','partyId','spawnId','runId','messageId','inputGuIds','guDef','guId','slotIDs','deadlock','foreignkey','datetime','sourceoftruth'])
LEX={x.lower() for x in LEX}
LEX.update('i ii iii trim casefold crit rng refund removed remove reroll uuid elapsed private economic commands unlock unlocks expired expiry snapshot carryover scope deadline baseline freeze readme chars outside ownerless reject rejected accepted accept all final exactly solo network frame sham standalone overhead default unknown notfound available capacity recipient source target param argument phi thì hấp phiên khắc bế tắc nghỉ lễ nhãn nét cách nén ước chốt ngưỡng phân hóa tăng dự toán giá bản minh họa'.split())
LEX.update(re.findall(r'(?<!\w)[A-Za-zÀ-ỹ]+(?!\w)',(OUT/'tools'/'plan_template.md').read_text(encoding='utf-8').split('## 6.')[0]))
LEX={x.lower() for x in LEX}
LEX.update(w+'s' for w in EN_WORDS)
LEX.update('identity primitive concurrent concurrency quota expiry opaque decimal multiple operator compound filename filesystem before after signature verified unverified readiness running deadlock versioned benchmark invite recipient payload length pixel ownership validation persistence dictionary projection readonly gameplay scope baseline resource artifact package packages leafpackage leaf packages money moderation foreign critical callback condition status skilldamage participation wipe grace rejoin internet source render renderer blocked pending config unix windows elapsed mode pattern effect primitive null sleep signing mapping stale rejected accepted ack pause locked rank_factor registry lookup outcome terminal newVersion firstValidHit sameID params php python casefold toArray eventID math source runbook sourceoftruth scan trust charge chance cls payload hash pressure label integer receipt durable immutable atomic schema flag region lean signing docs namespace path'.split())
LEX.update('gợi kiến trúc gửi trị bỏ đề phần lễ cập nhật cung chuyển quản số cấu rõ giảm nén chắc khoảng thành công dự phạt thủ chủ thể phạm vi hiệu lực hậu trạng thông vật giấy tái nạp phản ứng đầy đủ nhận định cộng đo cân suất nghiệm toàn xác phiên bản chức năng tiêu chí định lượng dừng bế tắc ước'.split())
LEX.update('domain rebaseline installer export port project commercial launch hardcore soft public native trusted main body completed complete payload bound transition duration stat stats global reset despawn pending chance output input primitives gameplay body ack applied sleep exception negatives invalid compatible explicitly tested sleep restart unavailable restore expired expiry expire participation deny classified playable foreign manifest reference styleguide depthsort decisionlog bottleneck quiz cutline sound owner prefix path testcase readonly tutorial macro publish progression consumed seal metadata option stage network sleep free lifetime online foreground background room worker quota CPU RAM memory payload'.lower().split())
LEX.update('median percentile artist social payment desktop multi swept collision dispel debuff bounded rulebook issue issues knowledge remaining implemented technical workflow framework realpayment kickoff copy word markdown openpyxl installer secret skills generate source compatible leaf packages fps loss important unused'.split())
LEX.update('coi chắn nhất quán mai mạng ích nghỉ ngày lễ buổi kích cỡ đào bào họng nén xuất chuyện nguồn phê duyệt phức hướng đo mật nghĩa đồng thuận sát lượng trận chiếm phía quy tắc tiền ngân sách giọng phím gõ định cứng nghìn thừa hết đủ tập mức thư mục khôi phục huyền bản mã căn bậc tài nguyên số lượng tông môn đồ sát phẩm chất'.split())
# Recover ordinary Vietnamese syllables from the naturally spaced source prose.
# This also protects short words from accidental decomposition.
vowels='aeiouyàáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹ'
syllable=re.compile(r'(?:ngh|ng|gh|ch|nh|th|tr|ph|kh|gi|qu|[bcdđghklmnpqrstvx])?['+vowels+r']+(?:ch|nh|ng|c|m|n|p|t)?$')
corpus=(OUT/'tools'/'plan_template.md').read_text(encoding='utf-8')+(OUT/'tools'/'planning_data.py').read_text(encoding='utf-8')
for token in re.findall(r'(?<!\w)[A-Za-zÀ-ỹ]+(?!\w)',corpus):
    tones=sum(ch in '\u0300\u0301\u0309\u0303\u0323' for ch in unicodedata.normalize('NFD',token))
    if len(token)<=8 and tones<=1 and syllable.fullmatch(token.lower()):LEX.add(token.lower())

REPLACE = {
 'singlewriter':'single writer','safespawn':'safe spawn','fulltime':'full time',
 'coreloop':'core loop','taskcarryover':'task carryover','frameTime':'frameTime',
 'frametime':'frame time','faulttests':'fault tests','maproom':'map room',
 'techoptions':'tech options','successspawnremoved':'success: spawn removed',
 'outputbound':'output bound','unlockrank-up':'unlock rank-up',
 'requestUUID':'request UUID','RNGdamage':'RNG damage','PKdrop':'PK drop',
 'ACKkhông':'ACK không','khóaeconomiccommands':'khóa economic commands',
 'claimunique':'claim unique','configVersion':'configVersion',
 'lowgraphics':'low graphics','stalelease':'stale lease','tickoverrun':'tick overrun',
 'oneinoneout':'one in, one out','foreignkeys':'foreign keys',
 'failCount':'failCount','QAt racker':'QA tracker',
 'tài nguyễnsẵncó':'tài nguyên sẵn có','tài nguyễnsẵn':'tài nguyên sẵn',
 'rankII':'rank II','rankI':'rank I',
 'CổII':'Cổ II','CổI':'Cổ I','CổChânNhân':'Cổ Chân Nhân',
 'knownissues':'known issues','decisionlog':'decision log','depthsort':'depth sort',
 'styleguide':'style guide','firstplayable':'first playable','lowgraphics':'low graphics',
 'trướcG1':'trước G1','trướcG3':'trước G3','đểP2':'để P2',
 'TrướcS16':'Trước S16','SauS01':'Sau S01','cuốiS01':'cuối S01','cuốiS02':'cuối S02',
 'cócorrelationID':'có correlationID','Questreward':'Quest reward','Terminalstatewriteonce':'Terminal state write once',
 'battleterminaluniquedeathID':'battle terminal unique deathID',
 'economiccommitbeforeACK':'economic commit before ACK','Explicitlyrejectmissingtoken':'Explicitly reject missing token',
 'capturabletargetlowHP':'capturable target low HP','rollbackcompatibleapp':'rollback compatible app',
 'coreloopE2E':'core loop E2E','chatXSS':'chat XSS','duelDC':'duel DC','restoreDB':'restore DB',
 'gọiGMgrant':'gọi GM grant','Mục tiêuNFR':'Mục tiêu NFR','theoNFR':'theo NFR',
 'medianFPS':'median FPS','ACKp95':'ACK p95','tickp95':'tick p95','timep95':'time p95',
 'vớiRTT':'với RTT','disasterrestoreRPO':'disaster restore RPO',
 '12luồngE2E':'12 luồng E2E','có79dòng':'có 79 dòng','có72leafpackage':'có 72 leaf packages',
 'dưới6deliverable':'dưới 6 deliverable','cần0P0':'cần 0 P0',
 'mạng20Mbps':'mạng 20 Mbps','36tuầnở20h':'36 tuần ở 20h',
 '20CổđầyđủtrướcG3':'20 Cổ đầy đủ trước G3',
 '08tổng4500XP':'08 tổng 4500 XP','khôngcopyAPIkhácmajor':'không copy API khác major',
 'ghiNFRunverifiedvàchốtlại':'ghi NFR unverified và chốt lại',
 'UHD620hoặcmáythực':'UHD620 hoặc máy thực',
 'demoPC':'demo PC','DemoPC':'Demo PC','Scope/FSDv0.1':'Scope/FSD v0.1',
 'gateG0':'gate G0','danh sáchTBD':'danh sách TBD',
 'ngânsách0':'ngân sách 0','Ngânsách0':'Ngân sách 0',
 '5Cổfirstplayable':'5 Cổ first playable','20Cổ':'20 Cổ',
 'Capturecomplete':'Capture complete','outputII':'output II','Guunique':'Gu unique',
 'CCduration':'CC duration','rankfactor':'rankFactor','chỉterminaltransition':'chỉ terminal transition',
 'deadlockretry':'deadlock retry','acceptedcast':'accepted cast','rejectedcast':'rejected cast',
 'unlockrank':'unlock rank','newroom':'new room','wincondition':'win condition',
 'và10MAT':'và 10 MAT','+3MAT':'+3 MAT','+2MAT':'+2 MAT',
 '7GM':'7 GM','và12luồng':'và 12 luồng','giờhòa':'giờ hòa',
 'sốngày':'số ngày','ngânsách':'ngân sách','bảnđầu':'bản đầu',
 'vàtelegraphs':'và telegraphs','mắtthươngmại':'mắt thương mại',
 'Khôngpayment':'Không payment','Khôngrealpayment':'Không real payment',
 'quảntrịvàthiếtkếeconomy':'quản trị và thiết kế economy',
 'Đónggóiclientdesktop':'Đóng gói client desktop',
 'thiếuartisthoặcbackendthìphảigiảmcontent':'thiếu artist hoặc backend thì phải giảm content',
 'Cóthểcắtsocialtrước':'Có thể cắt social trước','Redisđểsaukhimulti':'Redis để sau khi multi',
 'Trướckickoff':'Trước kickoff','vàsweptcollision':'và swept collision',
 'dispelgỡdebuff':'dispel gỡ debuff','retry3lầnbounded':'retry 3 lần bounded',
 'rankFactorảnhhưởngdamagekhôngnhânmọistat':'rankFactor ảnh hưởng damage, không nhân mọi stat',
 'vàescaperender':'và escape render','30sthualuật':'30s thua theo luật',
 'modelđềxuất':'model đề xuất','migrationchốtS02':'migration chốt S02',
 'khôngchứasecret':'không chứa secret','côngthứcrulebook':'công thức rulebook',
 'FSgiữasprint':'FS giữa sprint','dừngsocial':'dừng social','fixdữliệutrước':'fix dữ liệu trước',
 'đểghiissues':'để ghi issues','completevàlast':'complete và last',
 'participationruleđúng':'participation rule đúng','doubleDChòa':'double DC hòa',
 'khôngđụngDBnguồn':'không đụng DB nguồn','dùngscenario10':'dùng scenario 10',
 'khôngchặnwarphoặc2socketgâydupe':'không chặn warp hoặc 2 socket gây dupe',
 'vàpeerreview':'và peer review','transactiontrướcsocial':'transaction trước social',
 'sức1artist':'sức 1 artist','phảighiNFRthayđổi':'phải ghi NFR thay đổi',
 'sánhfunnel':'sánh funnel','knowledgechỉ1':'knowledge chỉ 1',
 'ôremainingâmphảicắtphạmvihoặclậplạilịch':'ô remaining âm phải cắt phạm vi hoặc lập lại lịch',
 '36tuầnlàphươngán20h':'36 tuần là phương án 20h','0làplaceholder':'0 là placeholder',
 'cầnnghiêncứugiàthịtrườngđểdùngbảngnày':'cần nghiên cứu giá thị trường để dùng bảng này',
 'thayđơngiáđểướcgiátrịcôngsứcnếumuốn':'thay đơn giá để ước giá trị công sức nếu muốn',
 'khôngthuêmớingườilàm':'không thuê mới người làm','dấuimplemented':'dấu implemented',
 'Technicaldocsđối':'Technical docs đối','khôngmangworkflowB2B':'không mang workflow B2B',
 'khôngkhẳngđịnh3':'không khẳng định 3','Khôngmặcđịnhframeworktựchốnghack':'Không mặc định framework tự chống hack',
 'hãylưubảncopyvàđưathayđổivềsource':'hãy lưu bản copy và đưa thay đổi về source',
 'cácNFRchưađượcbenchmark':'các NFR chưa được benchmark','vớiopenpyxl':'với openpyxl',
 'vàDoD':'và DoD','TrongExcel':'Trong Excel','bảnWord':'bản Word',
 'bảnMarkdown':'bản Markdown','củaWord':'của Word',
}

@lru_cache(maxsize=None)
def split_token(token):
    lower=token.lower()
    if lower in LEX or len(token)<4 or token.isupper() or token.isdigit(): return token
    if token.isascii() and re.search(r'[a-z][A-Z]',token):return token
    # Only accept a complete dictionary decomposition; leave unfamiliar names/identifiers alone.
    n=len(token); best={0:(0,[])}
    for i in range(n):
        if i not in best: continue
        for j in range(i+1,min(n,i+22)+1):
            part=lower[i:j]
            if part in LEX or part.isdigit():
                score=best[i][0] + len(part)**2 - 2
                if j not in best or score>best[j][0]: best[j]=(score,best[i][1]+[token[i:j]])
    if n not in best:return token
    parts=best[n][1]
    if token.isascii() and any(len(p)<4 and not p.isdigit() for p in parts):return token
    return ' '.join(parts)

def clean(value):
    if not isinstance(value,str) or '://' in value: return value
    for old,new in REPLACE.items():value=value.replace(old,new)
    value=re.sub(r'[^\W_]+',lambda m:split_token(m.group()),value,flags=re.UNICODE)
    value=re.sub(r'([,;:%])(?=[A-Za-zÀ-ỹ])',r'\1 ',value)
    value=re.sub(r'\b(rank [IVX]+)(?=\d)',r'\1 ',value)
    value=re.sub(r'\b(level|rank|hunger|channel|range|stack|buy|sell|owner|party|seal|slot|skill|fee|maxMana|maxHP|XP|XPcum|TTL|renew|countdown|trận|Số|NPC)(?=\d)',r'\1 ',value)
    value=re.sub(r'(\d)(?=(?:giây|giờ|phút|tuần|người|lượt|điểm|ký|tile|mana|currency|player|mob|tester|client|socket|Gu|gu|material|receipt|request|sprint|entry|reward|callback|writer|effect|output|input|slot|message|hunger|build|event|quest|map|Cổ|công|lần|hướng|pattern|primitive|snapshot|member|party|room|char|second|daily|weekly|NPC|PD|XP|VFX|NFR|testcase)\b)',r'\1 ',value)
    value=re.sub(r'(?<=[A-Za-zÀ-ỹ])(?=\d+(?:giây|giờ|phút|tuần|người|lượt|tile|mana|currency|player|mob|tester|client|Gu|hunger|daily|weekly)\b)', ' ',value)
    return re.sub(' +',' ',value).strip()

def cleaned(row): return [clean(v) for v in row]

TASKS=[]; phase_count=defaultdict(int)
for sn,sprint in enumerate(d.SPRINTS,1):
    parent,title=sprint[:2]
    for i,role in enumerate(['GD','CL','SV','AR']):
        name,dod,deps,effort=sprint[i+2]; phase_count[parent]+=1
        TASKS.append(dict(id=f'T{sn:02}-{role}',code=parent.split('.')[0]+'.'+str(phase_count[parent]),
          parent=parent,phase=title,name=name,dod=dod,role=role,sprint=sn,
          sw=sn*2-1,ew=sn*2,deps=deps,effort=effort))
TASK_BY_ID={t['id']:t for t in TASKS}
REQ_BY_ID={r['id']:r for r in d.REQS}
BASE=sum(t['effort'] for t in TASKS)
TOTAL=BASE*(1+d.META['reserve'])
CAP=4*d.META['weeks']*d.META['hours']*d.META['focus']/8
ROLE_BASE={r:sum(t['effort'] for t in TASKS if t['role']==r) for r,_,_ in d.ROLES}
def refs_for_task(id): return ', '.join(r['id'] for r in d.REQS if id in r['wbs'].split(','))

wb=Workbook(); wb.remove(wb.active)
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
wb.properties.title='GuMaster2D — Project Plan / FSD / WBS'
wb.properties.subject='Đề xuất closed alpha cho đội 4 người, Godot PC Windows'
wb.properties.creator='GuMaster2D planning'
CACHE=defaultdict(dict)

def formula(ws,coordinate,expression,cached):
    ws[coordinate]=expression;CACHE[ws.title][coordinate]=cached

def sheet(name,title,subtitle,headers,rows,widths=None):
    ws=wb.create_sheet(name); n=len(headers)
    ws.sheet_view.showGridLines=False
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
    ws.cell(1,1,clean(title));ws.cell(1,1).font=Font(name='Calibri',size=19,bold=True,color='FFFFFF')
    ws.cell(1,1).fill=PatternFill('solid',fgColor=NAVY);ws.row_dimensions[1].height=34
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n)
    ws.cell(2,1,clean(subtitle));ws.cell(2,1).font=Font(name='Calibri',size=10,italic=True,color=GRAY)
    ws.cell(2,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=31
    for j,h in enumerate(headers,1):
        c=ws.cell(4,j,h);c.font=Font(name='Calibri',size=10,bold=True,color='FFFFFF')
        c.fill=PatternFill('solid',fgColor=TEAL);c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[col(j)].width=(widths[j-1] if widths and j<=len(widths) else 25)
    ws.row_dimensions[4].height=31
    for rn,row in enumerate(rows,5):
        max_lines=1
        for j,value in enumerate(cleaned(row),1):
            c=ws.cell(rn,j,value);c.font=Font(name='Calibri',size=10,color=INK)
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if rn%2==1:c.fill=PatternFill('solid',fgColor='F0F5F8')
            c.border=Border(bottom=Side(style='hair',color='D9E3EB'))
            if isinstance(value,str):
                w=ws.column_dimensions[col(j)].width
                max_lines=max(max_lines,sum(max(1,math.ceil(len(x)/(w*0.95))) for x in value.split('\n')))
        ws.row_dimensions[rn].height=min(130,max(31,15*max_lines+9))
    ws.freeze_panes='C5';ws.auto_filter.ref=f'A4:{col(n)}{max(5,ws.max_row)}'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:4';ws.print_options.horizontalCentered=True
    ws.oddFooter.center.text='GuMaster2D | Đề xuất v0.2 | &P / &N'
    ws.sheet_properties.outlinePr.summaryRight=False
    return ws

def dropdown(ws,range_,items):
    dv=DataValidation(type='list',formula1='"'+','.join(items)+'"');dv.errorTitle='Giá trị chưa hợp lệ'
    dv.error='Chọn một giá trị trong danh sách.';dv.showErrorMessage=True
    ws.add_data_validation(dv);dv.add(range_)

ws=sheet('00_START','GuMaster2D / PLAN • FSD • WBS',
 '07/10/2026 • Bản đề xuất cho đội 4 người tự phát triển • Các ô nền vàng là giả định có thể sửa; ngày khởi động chỉ minh họa.',
 ['Tham số / Chỉ số','Giá trị','Diễn giải'],[],[38,28,110])
params=[
 (5,'Phiên bản',d.META['version'],'Chưa phải baseline đã được đội phê duyệt; có prototype Godot local; chưa nghiệm thu alpha. Xem docs/implementation và quality-report.'),
 (6,'Số thành viên',4,'Bốn vai trò kiêm nhiệm: GD/QA, Client, Server/Tech Lead, 2D/Content.'),
 (7,'Ngày bắt đầu minh họa',d.META['start'],'Thứ Hai 12/10/2026; sửa để dịch lịch. Khung ngày theo thứ Hai–thứ Sáu, chưa trừ lễ/nghỉ.'),
 (8,'Giờ / người / tuần',20,'Giả định vì user chưa cung cấp availability. Sửa giờ làm thay đổi công suất, không tự nén 18 sprint.'),
 (9,'Tỷ lệ thời gian hữu ích',0.75,'25% cho trao đổi, học, review, hỗ trợ; effort WBS là giờ hữu ích.'),
 (10,'Dự phòng effort',0.20,'20% trên effort cơ sở; đã tách riêng, không phải thêm 20% tuần sau tuần36.'),
 (11,'Số tuần kế hoạch',36,'18 sprint ×2tuần. Nếu đổi số sprint/phạm vi, cần lập lại WBS và Gantt.'),
 (12,'Giờ / person-day',8,'1PD=8giờ hữu ích; effort khác thời gian lịch và khác một ngày làm bán thời gian.'),
 (13,'Effort cơ sở (PD)',None,'Tổng72gói công việc, chưa dự phòng.'),
 (14,'Effort gồm dự phòng (PD)',None,'Cơ sở ×(1+dựphòng).'),
 (15,'Giờ hữu ích cần có',None,'Effort gồm dựphòng×8.'),
 (16,'Công suất đội (PD)',None,'Sốngười×tuần×giờ/tuần×tỷlệhữuích÷8.'),
 (17,'Công suất còn lại (PD)',None,'Côngsuất−effort gồm dựphòng. Âm thì cần giảmscope hoặc tănglịch.'),
 (18,'Yêu cầu chức năng',len(d.REQS),'30Client +18Server +7Admin; mỗi ID có AC, WBS và test.'),
 (19,'Yêu cầu phi chức năng',len(d.NFR),'Các mục tiêu NFR cần đo; chưa cóbenchmark thực tế.'),
 (20,'WBS leaf packages',len(TASKS),'Có code phâncấp,owner,predecessor,effort,lịch,DoD vàFSD links.'),
 (21,'Mốc kết thúc minh họa',None,'G5:cuối tuần36; closedalpha, chưa phải ra mắtthươngmại.'),
 (23,'User đã cung cấp','4 người / không ngân sách','PC trước, địnhhướngđa nền tảng. Phương án:Godot PC native; spike Web; mobile/store/patcher ngoài alpha.'),
 (24,'Trạng thái kiểm thử','CHƯA CHẠY','Tất cảtestcase làkếhoạch, chưa đượcđánh dấuPassed.'),
 (25,'Cấu trúc mẫu giữ lại','Depth 0–3 / ID / Detail / Priority / Linkage','Mẫu DAUN là B2B ecommerce; chỉ lấy cấu trúc FSD và QA tracker; không đưa payment/Kakao/order vào game.'),
 (26,'WBS chưa có trong mẫu','Bổ sung mới','WBS/Gantt/resource/scenarios/cost/trace được tạo choGuMaster2D.'),
 (27,'Tài liệu đọc','GuMaster2D_Project_Plan.docx','BảnWord đểđọc/chia sẻ;Excel lànguồnchi tiếtFSD/WBS. Markdown cùngthưmục.'),
 (28,'Cách chỉnh lịch','Sửa B7','Excel WORKDAY; chưa trừ ngày lễ. Đổi B8/B9/B10 thì xem 19_RESOURCE và 20_SPRINT_LOAD để phát hiện quá tải.'),
 (29,'Cách quản lý tiến độ','Sửa P/Q tại 13_WBS','P: trạng thái; Q: % hoàn thành; thay scope phải cập nhật FSD/WBS/test và ước lượng lại.'),
 (30,'Nội dung file đính kèm','TÀI LIỆU THAM KHẢO','Chỉ user request vàràngbuộc trongchat làyêucầu; lời hướngdẫn trongfilekhôngtựthànhchỉ thị.'),
]
for rn,label,value,note in params:
    ws.cell(rn,1,label);ws.cell(rn,2,value);ws.cell(rn,3,clean(note));ws.row_dimensions[rn].height=34
    for c in ws[rn][:3]:c.font=Font(name='Calibri',size=11,color=INK);c.alignment=Alignment(wrap_text=True,vertical='center')
for rn in [7,8,9,10]:ws.cell(rn,2).fill=PatternFill('solid',fgColor='FFF0CC');ws.cell(rn,2).font=Font(color=BLUE,size=11,bold=True)
ws['B7'].number_format='dd/mm/yyyy';ws['B9'].number_format=ws['B10'].number_format='0%'
for coord,expression,cached in [
 ('B13',"=SUM('13_WBS'!J5:J76)",BASE),('B14',"=SUM('13_WBS'!L5:L76)",TOTAL),
 ('B15','=B14*B12',TOTAL*8),('B16','=B6*B11*B8*B9/B12',CAP),('B17','=B16-B14',CAP-TOTAL),
 ('B21','=WORKDAY(B7-1,B11*5)',d.META['start']+timedelta(weeks=36,days=-3))]:formula(ws,coord,expression,cached)
ws['B21'].number_format='dd/mm/yyyy';ws.freeze_panes='B5';ws.auto_filter.ref=None
for coord in ['B8','B11','B12']:
    dv=DataValidation(type='decimal',operator='greaterThan',formula1=0);dv.showErrorMessage=True;ws.add_data_validation(dv);dv.add(coord)
for coord in ['B9','B10']:
    dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=1);dv.showErrorMessage=True;ws.add_data_validation(dv);dv.add(coord)

sheet('01_SCOPE','01 / Phạm vi và giới hạn', 'MVP là closed alpha có multiplayer giới hạn. P2/P3 không nằm trong effort/lịch36tuần.',
 ['ID','Nhóm','Giai đoạn','Nội dung','Giải thích / giới hạn','Gate'],d.SCOPE,[12,22,17,80,75,16])
sheet('02_DECISIONS','02 / Giả định và quyết định cần chốt', 'Ràng buộc user đã nêu được phân biệt với đề xuất thiết kế và giả định ước lượng.',
 ['ID','Quyết định','Trạng thái','Phương án cơ sở','Owner','Hạn chốt','Tác động'],d.DECISIONS,[12,26,21,80,23,23,80])

FSD_HEADERS=['No','0 Depth','1 Depth','2 Depth','3 Depth','Function / Requirement Name','Requirement ID','Detail / Quy tắc xử lý','Category','Importance','Linkage / Req liên quan','Note / Nguồn','Input / Tiền điều kiện','Output / Hậu điều kiện','Validation / Ngoại lệ','Acceptance Criteria','WBS IDs','Test case','Business Rules','Use Cases','Phase','Status']
for prefix,name,title in [('CL','03_FSD_CLIENT','03 / FSD — Player Client'),('SV','04_FSD_SERVER','04 / FSD — Server & Persistence'),('AD','05_FSD_ADMIN','05 / FSD — GM / Admin')]:
    selected=[r for r in d.REQS if r['id'].startswith(prefix)]
    rows=[]
    for i,r in enumerate(selected,1):
        linked=','.join(sorted({x['id'] for x in d.REQS if x['id']!=r['id'] and set(r['wbs'].split(','))&set(x['wbs'].split(','))}))
        rows.append([i,r['module'],r['group'],r['screen'],'Action',r['name'],r['id'],r['spec'],prefix,r['priority'],linked,'Đề xuất từ ý tưởng + cụ thể hóa cho alpha',r['inputs'],r['outputs'],r['errors'],r['ac'],r['wbs'],r['test'],r['rule'],r['uc'],r['phase'],r['status']])
    ws=sheet(name,title,'12cộtđầu tương ứngschemaDAUN; bổsunginput/output/errors/AC/WBS/test/rules/UC/phase/status. Không cóyêucầunào đãđượcđội duyệt.',FSD_HEADERS,rows,[6,18,19,19,13,35,17,80,13,15,35,38,45,45,60,75,31,20,24,20,15,18])
    ws.freeze_panes='H5';dropdown(ws,f'V5:V{ws.max_row}',['Đề xuất','Đã chốt','Đang làm','Hoàn thành','Hoãn'])
sheet('06_NFR','06 / Non-functional Requirements','Mỗi chỉ số là mục tiêu đo trong điều kiện ghi rõ; chưa kiểm chứng bằng phần mềm thực tế.',
 ['ID','Nhóm','Điều kiện đo','Tiêu chí pass','WBS IDs','Evidence cần lưu'],d.NFR,[15,25,90,90,35,70])
sheet('07_RULES','07 / Business Rules','Thông sốv0.1đểlập trình/thửnghiệm;chưalàcânbằngđãplaytest. Thayđổi phảităngconfigVersion.',
 ['ID','Nhóm','Quy tắc / Công thức','Biên / Ngoại lệ','Requirement IDs'],d.RULES,[15,25,110,95,55])
sheet('08_USE_CASES','08 / Luồng sử dụng','Happy path + lỗi + dữ liệu thay đổi. Các thao tác tài sản phải transaction và idempotent.',
 ['ID','Actor','Tiền điều kiện','Luồng chính','Ngoại lệ / Phục hồi','Dữ liệu liên quan','Requirements'],d.USE_CASES,[15,25,65,100,100,55,55])
sheet('09_DATA','09 / Data Dictionary','Logical modelđềxuất;chi tiếttype/index/migrationchốtS02. PostgreSQL lànguồnchotàisản;roommemory chothếgiớirealtime.',
 ['Entity','Fields / Keys','Relations / Constraints','Owner','Lưu ý'],d.DATA,[27,100,75,16,100])
sheet('10_API','10 / API & Realtime Contract','Têncommand/endpoint ởmứcthiếtkế,khôngcamkếtAPIframework. Validate auth/ownership/schema/rate trênserver.',
 ['ID','Transport','Endpoint / Command','Quyền','Request','Response','Error codes','Requirement IDs'],d.API,[15,16,60,35,75,70,65,45])
ws=sheet('11_GU_CATALOG','11 / Catalog 20 Cổ','10CổrankI+10rankII;tênminhhọa;nhiềubiếnthểdùng8effectprimitives. power củaheal/shield theo côngthứcrulebook;buff cóđơn vịriêng.',
 ['ID','Tên','Rank','Loại','Primitive','Target','Base power / effect value','Mana %max','CD (s)','Range (tile)','Duration (s)','Food ID','Nguồn / Ghi chú effect','Trạng thái cân bằng'],d.GU,[14,28,10,20,18,18,24,16,14,18,18,16,65,55])
ws.freeze_panes='D5'
sheet('12_CONTENT','12 / Maps • Mobs • Items • Recipes • Quests','Catalog nhápv0.1;giá/sốlượnglàgiảđịnhthiếtkế,khôngphảidữliệuthịtrường. Q01–08tổng4500XP,đủlevel10trướcrankII.',
 ['Loại','ID','Tên / Inputs','Chi tiết / Luật','Điều kiện / Output','Owner / Reward'],d.CONTENT+
 [('Recipe',r[0],r[1],f'{r[2]}; {r[3]}; fee {r[4]}; chance {r[6]}% +{r[7]}%/fail',r[5],r[8]) for r in d.RECIPES]+
 [('Quest '+q[1],q[0],q[2],q[4],q[3],f'{q[5]} XP; {q[6]} currency; {q[7]}') for q in d.QUESTS],
 [22,18,42,105,75,80])

rows=[]
PROTOTYPE_TASKS = {'T01-GD','T01-CL','T01-AR'}
for t in TASKS:
    rows.append([t['code'],t['id'],t['phase'],t['name'],t['role'],t['sprint'],t['sw'],t['ew'],t['deps'],t['effort'],None,None,None,None,None,'Đang làm' if t['id'] in PROTOTYPE_TASKS else 'Chưa bắt đầu',0,t['dod'],refs_for_task(t['id'])])
ws=sheet('13_WBS','13 / Work Breakdown Structure — 72 work packages','Effort=person-day8giờhữuích;deadlinekhung2tuần;FSgiữasprint,phốihợp trongcùngsprint. Dựphòng đãcótrong36tuần.',
 ['WBS code','Task ID','Deliverable group','Work package','Owner','Sprint','Start week','End week','Predecessors (FS)','Base PD','Reserve PD','Total PD','Hours incl reserve','Start date','End date','Status','Done %','Definition of Done','Requirement IDs'],rows,
 [14,17,32,56,12,12,14,14,44,15,15,15,20,18,18,21,13,85,60])
for rn,t in enumerate(TASKS,5):
    formula(ws,f'K{rn}',f"=J{rn}*'00_START'!$B$10",t['effort']*.2)
    formula(ws,f'L{rn}',f'=J{rn}+K{rn}',t['effort']*1.2)
    formula(ws,f'M{rn}',f"=L{rn}*'00_START'!$B$12",t['effort']*1.2*8)
    formula(ws,f'N{rn}',f"=WORKDAY('00_START'!$B$7-1,(G{rn}-1)*5+1)",d.META['start']+timedelta(weeks=t['sw']-1))
    formula(ws,f'O{rn}',f"=WORKDAY('00_START'!$B$7-1,H{rn}*5)",d.META['start']+timedelta(weeks=t['ew'],days=-3))
    ws[f'N{rn}'].number_format=ws[f'O{rn}'].number_format='dd/mm/yyyy'
    for c in ['J','K','L','M']:ws[f'{c}{rn}'].number_format='0.0'
    ws[f'Q{rn}'].number_format='0%';ws[f'J{rn}'].font=Font(color=BLUE,size=10)
dropdown(ws,'P5:P76',['Chưa bắt đầu','Đang làm','Đang review','Hoàn thành','Bị chặn','Hoãn'])
dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=1);dv.showErrorMessage=True;ws.add_data_validation(dv);dv.add('Q5:Q76')
ws.freeze_panes='E5'

rows=[[t['id'],t['name'],t['role'],t['sw'],t['ew']]+[None]*36 for t in TASKS]
ws=sheet('14_GANTT','14 / Lịch theo tuần','Ômàuxanh=khungthựchiện;khôngphảisốngàyfulltime. Mỗiowner1package/sprint vàreviewchéo từ25%thờigian khônghữuích.',
 ['Task ID','Work package','Owner','Start week','End week']+list(range(1,37)),rows,[17,55,12,12,12]+[4]*36)
for rn,t in enumerate(TASKS,5):
    # Gantt responds to Start/End week changes on the matching WBS row.
    formula(ws,f'D{rn}',f"='13_WBS'!G{rn}",t['sw']);formula(ws,f'E{rn}',f"='13_WBS'!H{rn}",t['ew'])
    for j,week in enumerate(range(1,37),6):
        formula(ws,f'{col(j)}{rn}',f'=IF(AND({col(j)}$4>=$D{rn},{col(j)}$4<=$E{rn}),1,0)',int(t['sw']<=week<=t['ew']))
        ws.cell(rn,j).number_format=';;;';ws.cell(rn,j).alignment=Alignment(horizontal='center')
    ws.row_dimensions[rn].height=30
ws.conditional_formatting.add('F5:AO76',CellIsRule(operator='equal',formula=['1'],fill=PatternFill('solid',fgColor=TEAL)))
ws.freeze_panes='F5'

rows=[]
for gid,sn,name,accept,owner,fail in d.GATES:
    rows.append([gid,sn,name,None,accept,owner,fail,'Chưa đánh giá',''])
ws=sheet('15_GATES','15 / Milestones & Exit Criteria','Gatechỉpasskhicóevidence;thấtbại→sửdụngdựphòng/cắtphạmvi/rebaseline. Mốckết thúcminhhọađượcdịchtừ00_START!B7.',
 ['Gate','End sprint','Mốc','Ngày minh họa','Exit criteria','Sign-off owner','Nếu chưa đạt','Status','Evidence / Notes'],rows,[12,14,42,19,105,30,95,22,60])
for rn,g in enumerate(d.GATES,5):
    formula(ws,f'D{rn}',f"=WORKDAY('00_START'!$B$7-1,B{rn}*10)",d.META['start']+timedelta(weeks=g[1]*2,days=-3));ws[f'D{rn}'].number_format='dd/mm/yyyy'
dropdown(ws,f'H5:H{ws.max_row}',['Chưa đánh giá','Đạt','Không đạt','Đạt có điều kiện'])

TEST_ROWS=[]
for r in d.REQS:
    last_sprint=max(int(x[1:3]) for x in r['wbs'].split(','))
    gate='G1' if last_sprint<=4 else 'G2' if last_sprint<=6 else 'G3' if last_sprint<=12 else 'G4' if last_sprint<=15 else 'G5'
    TEST_ROWS.append([r['test'],r['id'],r['name'],r['inputs'],r['spec']+' | Thử negative case: '+r['errors'],r['ac'],'P1',gate,'GD + owner kỹ thuật','Chưa chạy','',''])
for n in d.NFR:TEST_ROWS.append(['TC-'+n[0],n[0],n[1],n[2],n[5],n[3],'P1','G5','GD + CL/SV','Chưa chạy','',''])
for x in d.E2E_TESTS:TEST_ROWS.append([x[0],x[6],x[1],x[2],x[3],x[4],'P0' if x[0] in ['E2E-02','E2E-03','E2E-04','E2E-07'] else 'P1',x[5],'GD + CL/SV','Chưa chạy','',''])
ws=sheet('16_TEST_CASES','16 / Test Plan & Requirement Trace','79plannedtestrows;khôngtựđánh dấuPassed. Evidencephảicóbuild/config/browser/hardware khiápdụng; dùngQA_REQUESTS đểghiissues.',
 ['Test ID','Requirement IDs','Scenario','Preconditions / Input','Steps / Negative cases','Expected / Acceptance','Priority','Gate','Owner','Status','Evidence path','Actual result'],TEST_ROWS,[23,48,52,70,110,100,12,12,24,20,60,80])
dropdown(ws,f'J5:J{ws.max_row}',['Chưa chạy','Passed','Failed','Blocked','Not applicable'])

sheet('17_QA_REQUESTS','17 / QA Issue Tracker','ThamkhảosheetQAFromClientcủamẫu;đâylàformrỗng,khôngphảilỗipháthiệnthực tế. P0 mất/nhândữliệu, P1blocker, P2degraded, P3cosmetic.',
 ['Issue ID','Category','Function / Req ID','Path / Screen','Steps / Request','Requested date','Reporter','Status','Planned date','Completed date','Owner','Priority','Notes / Evidence'],[['']*13 for _ in range(12)],
 [17,22,25,28,100,19,23,20,19,19,18,13,70])
sheet('18_RISKS','18 / Risk Register','Likelihood/Impact1–5;score=L×I. Rủirocaođượcxem lạiởmỗigate;khôngcoidựphònglàcôngsuấtchothêmfeature.',
 ['ID','Risk','Likelihood','Impact','Trigger','Owner','Mitigation','Fallback / Cutline','Score'],[list(x)+[None] for x in d.RISKS],
 [14,48,14,14,90,20,100,95,14])
ws=wb['18_RISKS']
for rn,x in enumerate(d.RISKS,5):formula(ws,f'I{rn}',f'=C{rn}*D{rn}',x[2]*x[3])
ws.conditional_formatting.add(f'I5:I{ws.max_row}',ColorScaleRule(start_type='min',start_color=GREEN,end_type='max',end_color='E99887'))

rows=[]
for role,title,responsibilities in d.ROLES:
    rows.append([role,title,responsibilities,ROLE_BASE[role],None,None,None,None,None])
ws=sheet('19_RESOURCE','19 / Công suất và tải theo vai trò','Mỗi người20h/tuần×36tuần×75%=67.5PDhữuích. Dựphòng20%đãtrongtải;khônggiảđịnhcóQA/DevOps thứ5.',
 ['Role','Người / vai trò','Trách nhiệm','Base PD','Reserve PD','Total PD','Capacity PD','Remaining PD','Utilization'],rows,[12,52,100,17,17,17,18,18,18])
for rn,(role,_,_) in enumerate(d.ROLES,5):
    b=ROLE_BASE[role]
    formula(ws,f'D{rn}',f'=SUMIF(\'13_WBS\'!$E$5:$E$76,A{rn},\'13_WBS\'!$J$5:$J$76)',b)
    formula(ws,f'E{rn}',f"=D{rn}*'00_START'!$B$10",b*.2)
    formula(ws,f'F{rn}',f'=D{rn}+E{rn}',b*1.2)
    formula(ws,f'G{rn}',"='00_START'!$B$11*'00_START'!$B$8*'00_START'!$B$9/'00_START'!$B$12",CAP/4)
    formula(ws,f'H{rn}',f'=G{rn}-F{rn}',CAP/4-b*1.2)
    formula(ws,f'I{rn}',f'=IFERROR(F{rn}/G{rn},0)',b*1.2/(CAP/4));ws[f'I{rn}'].number_format='0%'
ws.conditional_formatting.add('H5:H8',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor=RED)))
chart=BarChart();chart.title='Tải gồm dự phòng / Công suất (PD)';chart.y_axis.title='PD';chart.x_axis.title='Vai trò'
chart.add_data(Reference(ws,min_col=6,max_col=7,min_row=4,max_row=8),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=5,max_row=8));chart.width=19;chart.height=9;ws.add_chart(chart,'A12')

rows=[[t['sprint'],t['role'],t['id'],None,None,None,None,None] for t in TASKS]
ws=sheet('20_SPRINT_LOAD','20 / Kiểm tra tải từng sprint × owner','Mỗi2tuần:capacity3.75PD;package3PD+20%=3.6PD(96%). Nếuđổigiờ/effort,ôremainingâmphảicắtphạmvihoặclậplạilịch.',
 ['Sprint','Owner','Baseline Task ID','Base PD','Total PD incl reserve','Capacity PD / sprint','Remaining PD','Check'],rows,[14,15,22,19,25,26,23,25])
for rn,t in enumerate(TASKS,5):
    wr=5+TASKS.index(t)
    formula(ws,f'D{rn}',f"=SUMIFS('13_WBS'!$J$5:$J$76,'13_WBS'!$E$5:$E$76,B{rn},'13_WBS'!$F$5:$F$76,A{rn})",t['effort'])
    formula(ws,f'E{rn}',f"=SUMIFS('13_WBS'!$L$5:$L$76,'13_WBS'!$E$5:$E$76,B{rn},'13_WBS'!$F$5:$F$76,A{rn})",t['effort']*1.2)
    formula(ws,f'F{rn}',"=2*'00_START'!$B$8*'00_START'!$B$9/'00_START'!$B$12",3.75)
    formula(ws,f'G{rn}',f'=F{rn}-E{rn}',3.75-t['effort']*1.2)
    formula(ws,f'H{rn}',f'=IF(G{rn}<0,"OVERLOAD","OK")','OK')
ws.conditional_formatting.add('G5:G76',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor=RED)))

rows=[]
for h in [10,15,20,30,40]:rows.append([h,None,None,'Quy đổi công suất; chỉ dùng khi cả4người cócùngavailability. Cầnrebaseline cácsprint/phụthuộc,khôngcoilàdeadlinechắcchắn.'])
ws=sheet('21_SCENARIOS','21 / Quy đổi thời gian theo availability','36tuầnlàphươngán20h/người/tuần; thờigian thực tếphụthuộcavailability/nghỉ/học vàkếtquảspike. BảngkhôngtựdịchWBS.',
 ['Hours / người / tuần','Equivalent weeks','Equivalent months (~4.345w)','Lưu ý'],rows,[27,26,31,110])
for rn,h in enumerate([10,15,20,30,40],5):
    formula(ws,f'B{rn}',f"='00_START'!$B$11*'00_START'!$B$8/A{rn}",36*20/h)
    formula(ws,f'C{rn}',f'=B{rn}/4.345',36*20/h/4.345)
    ws[f'B{rn}'].number_format='0.0';ws[f'C{rn}'].number_format='0.0'

cash=[
 ['CASH-01','Hosting server / tháng',0,0,0,'Chưa có báo giá / chưa có host','Local/LAN trước; điền đơn giá thật khi chọn môi trường.'],
 ['CASH-02','Backup storage / tháng',0,0,0,'Chưa có báo giá','Backup cần ở nơi độc lập với máy chạy; dữ liệualpha cóthểdùngtài nguyễnsẵncónếuđủđiềukiện.'],
 ['CASH-03','Domain / năm (tùy chọn)',0,0,0,'Chưa chọn','KhôngbắtbuộcchoalphaLAN/private; khôngmặcđịnhmiễnphí.'],
 ['CASH-04','Art/audio mua một lần',0,0,0,'Mục tiêu cash0','Tựtạohoặctài nguyêncónguồnsửdụngrõ;giới hạncontent đểkhôngcầnmua.'],
 ['CASH-05','Công cụ / tháng',0,0,0,'Chưa chọn','Dùngcôngcụsẵncó;mọikhoảntrảphíphảiđượcownerchốt.'],
 ['CASH-06','Tổng cash nhập',None,None,None,'Số0khôngphảidựtoánthịtrường','Tổngchỉcótácdụngsaukhiđiềncácđơngiá/sốlượng;khôngcamkếtonline0đồng.'],
 ['TIME-01','Giờ công cơ sở',BASE*8,0,None,'Giờ đóng góp, không phải lương','Đơngiá0theobốicảnhtựpháttriển; thayđơngiáđểướcgiátrịcôngsứcnếumuốn.'],
 ['TIME-02','Giờ công dự phòng',BASE*.2*8,0,None,'20% effort','Táchkhỏicashhạtầng;khôngthuêmớingườilàm.'],
]
ws=sheet('22_COST','22 / Cash input & giá trị công sức','User:no budget. Tấtcảđơngiá=0làplaceholder/mụctiêu,khôngphảibáogiáđãxácnhận. Không cầnnghiêncứugiàthịtrườngđểdùngbảngnày.',
 ['ID','Khoản','Số lượng / Giờ','Đơn giá VND (input)','Thành tiền VND','Trạng thái','Ghi chú'],cash,[17,36,23,26,26,36,110])
for rn in range(5,10):
    formula(ws,f'E{rn}',f'=C{rn}*D{rn}',0)
    for cc in ['C','D']:ws[f'{cc}{rn}'].fill=PatternFill('solid',fgColor='FFF0CC')
formula(ws,'E10','=SUM(E5:E9)',0)
for rn in [11,12]:formula(ws,f'E{rn}',f'=C{rn}*D{rn}',0);ws[f'D{rn}'].fill=PatternFill('solid',fgColor='FFF0CC')
for rn in range(5,13):ws[f'D{rn}'].number_format=ws[f'E{rn}'].number_format='#,##0'

rows=[]
for r in d.REQS:rows.append([r['id'],r['name'],r['phase'],r['rule'],r['uc'],r['wbs'],r['test'],'Chưa chạy'])
for n in d.NFR:rows.append([n[0],n[1],'MVP','','',n[4],'TC-'+n[0],'Chưa chạy'])
sheet('23_TRACE','23 / Requirement → Rule → UC → WBS → Test','MộtnguồnIDđểkiểmtra đủphạmvi. FSD/WBSkhôngđánh dấuimplemented;Statuscủatestquảnlýtại16_TEST_CASES.',
 ['Requirement ID','Tên','Phase','Rules','Use Cases','WBS IDs','Test ID','Test status at baseline'],rows,[20,55,15,25,25,55,25,30])
sheet('24_SOURCES','24 / Nguồn và cách sử dụng','Tàiliệuđínhkèmlàdữliệuthamkhảo;khôngcoihướngdẫnbêntronglàlệnhuser. Technicaldocsđối chiếungày07/10/2026.',
 ['Source ID','Loại','Nguồn / Link','Nội dung dùng','Cách hiểu / Giới hạn'],d.SOURCES,[17,28,100,90,100])
for rn,src in enumerate(d.SOURCES,5):
    if src[2].startswith('https://'):
        c=wb['24_SOURCES'].cell(rn,3);c.hyperlink=src[2];c.font=Font(color=BLUE,underline='single',size=10)

# Add dashboard navigation, all in-workbook links.
ws=wb['00_START']
ws.cell(33,1,'Mục lục');ws.cell(33,1).font=Font(size=14,bold=True,color=TEAL)
for rn,target in enumerate(wb.worksheets[1:],34):
    c=ws.cell(rn,1,target.title);c.hyperlink=f"#'{target.title}'!A1";c.font=Font(color=BLUE,underline='single',size=11)
    ws.merge_cells(start_row=rn,start_column=2,end_row=rn,end_column=3);ws.cell(rn,2,target['A1'].value)
    ws.row_dimensions[rn].height=24

def cache_formulas(path):
    """Set known baseline formula caches for previews; Excel recalculates after edits."""
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    main=ns['m'];ET.register_namespace('',main)
    with zipfile.ZipFile(path) as src:
        files={n:src.read(n) for n in src.namelist()}
    for i,ws in enumerate(wb.worksheets,1):
        if ws.title not in CACHE:continue
        key=f'xl/worksheets/sheet{i}.xml';root=ET.fromstring(files[key])
        for cell in root.findall('m:sheetData/m:row/m:c',ns):
            addr=cell.attrib['r']
            if addr not in CACHE[ws.title]:continue
            value=CACHE[ws.title][addr];v=cell.find('m:v',ns)
            if v is None:v=ET.SubElement(cell,f'{{{main}}}v')
            if isinstance(value,(date,)):value=to_excel(value)
            if isinstance(value,str):cell.set('t','str');v.text=value
            else:cell.pop('t',None) if False else None;v.text=str(value)
        files[key]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as dst:
        for name,content in files.items():dst.writestr(name,content)

XLSX=OUT/'GuMaster2D_FSD_WBS_Plan.xlsx'
wb.save(XLSX);cache_formulas(XLSX)

def validate():
    ids=set(TASK_BY_ID);reqids=set(REQ_BY_ID)|{n[0] for n in d.NFR}
    assert len(ids)==72 and len(REQ_BY_ID)==55
    for t in TASKS:
        for pre in filter(None,t['deps'].split(',')):
            assert pre in ids,(t['id'],pre)
            assert TASK_BY_ID[pre]['sprint']<t['sprint'],('FS overlap',pre,t['id'])
        assert t['effort']*1.2 <= 3.75+1e-8
    for r in d.REQS:
        assert all(x in ids for x in r['wbs'].split(',')),r['id']
        assert r['ac'] and r['inputs'] and r['outputs'] and r['errors'],r['id']
    for n in d.NFR:assert all(x in ids for x in n[4].split(',')),n[0]
    for x in TEST_ROWS:assert all(v.strip() in reqids for v in x[1].split(',')),x[:2]
    for r in d.RECIPES:
        assert all(g.strip() in {x[0] for x in d.GU} for g in r[1].split('+'))
        assert r[5] in {x[0] for x in d.GU}
    assert sum(q[5] for q in d.QUESTS[:8])==4500
    assert BASE==199 and abs(TOTAL-238.8)<1e-8 and CAP==270
    f=load_workbook(XLSX,data_only=False);v=load_workbook(XLSX,data_only=True)
    assert len(f.sheetnames)==25
    assert abs(v['00_START']['B14'].value-TOTAL)<1e-8
    assert v['00_START']['B21'].value.date()==date(2027,6,18)
    errors=[]; formulas=0
    for s in f:
        for row in s:
            for c in row:
                if c.data_type=='f':
                    formulas+=1;assert '#REF!' not in c.value
                    for ref in re.findall(r"'([^']+)'!",c.value):assert ref in f.sheetnames
                    assert v[s.title][c.coordinate].value is not None,(s.title,c.coordinate)
                if c.data_type=='e':errors.append((s.title,c.coordinate,c.value))
    assert not errors,errors
    assert all(v['20_SPRINT_LOAD'].cell(r,8).value=='OK' for r in range(5,77))
    return {'status':'PASS','scope':'Planning artifact integrity only; 79 formal alpha tests have NOT run. Godot prototype evidence is separate.',
      'sheets':len(f.sheetnames),'functional_requirements':55,'nfr':12,'wbs_packages':72,
      'test_rows':len(TEST_ROWS),'gu':20,'recipes':10,'quests':10,'formulas_with_baseline_cache':formulas,
      'effort_base_pd':BASE,'effort_with_reserve_pd':round(TOTAL,1),'capacity_pd':CAP,
      'capacity_by_role_pd':67.5,'effort_by_role_pd':ROLE_BASE,'overloaded_sprint_roles':0,
      'illustrative_start':'2026-10-12','illustrative_end':'2027-06-18'}

REPORT=validate()
(OUT/'validation_report.json').write_text(json.dumps(REPORT,ensure_ascii=False,indent=2),encoding='utf-8')

def md_table(headers,rows):
    def escape(x):return clean(str(x)).replace('|',' / ').replace('\n','<br>')
    return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+
      ['| '+' | '.join(escape(x) for x in row)+' |' for row in rows])

fsd_summary=[]
for label,prefix in [('Player Client','CL'),('Server & Persistence','SV'),('GM / Admin','AD')]:
    fsd_summary.append('### '+label+'\n\n'+md_table(['ID','Chức năng','Tiêu chí nghiệm thu tiêu biểu'],
       [[r['id'],r['name'],r['ac']] for r in d.REQS if r['id'].startswith(prefix)]))
phase_rows=[]
for parent,title in dict((s[0],s[1]) for s in d.SPRINTS).items():
    ts=[t for t in TASKS if t['parent']==parent];eff=sum(t['effort'] for t in ts)
    phase_rows.append([parent,title,len(ts),f'{eff:.1f}',f'{eff*.2:.1f}',f'{eff*1.2:.1f}'])
focuses=['Scope + spike','Luật + nền tảng','Điều khiển + session','Realtime + AOI','Combat5Cổ','PvE và loot','Bag / loadout / atomic economy','Capture / feed','Refine / pity / receipt','Quest / rank / NPC shop','20Cổ / content validator','Recovery / coreloopE2E','Party / chat','Dungeon / boss','Duel opt-in','GM / metrics / feature freeze','UAT / load / performance','RC / restore / alpha report']
roadmap=[]
for sn,sprint in enumerate(d.SPRINTS,1):
    gate=next((g[0] for g in d.GATES if g[1]==sn),'—')
    roadmap.append([f'S{sn:02} / W{sn*2-1}–{sn*2}',focuses[sn-1],sprint[3][1]+'; '+sprint[4][1],gate])
resrows=[[r,f'{ROLE_BASE[r]:.1f} / {ROLE_BASE[r]*1.2:.1f} PD','67.5 PD',f'{67.5-ROLE_BASE[r]*1.2:.1f} PD'] for r,_,_ in d.ROLES]
text=(OUT/'tools'/'plan_template.md').read_text(encoding='utf-8')
for key,value in {
 'FSD_SUMMARY':'\n\n'.join(fsd_summary),
 'PHASE_SUMMARY':md_table(['Nhóm','Deliverable','Packages','Base PD','Reserve PD','Total PD'],phase_rows),
 'SPRINT_ROADMAP':md_table(['Sprint / tuần','Trọng tâm','Đầu ra để review','Gate'],roadmap),
 'RESOURCE_SUMMARY':md_table(['Vai trò','Cơ sở / Gồm dự phòng','Công suất','Còn lại'],resrows),
}.items():text=text.replace('{{'+key+'}}',value)
text='\n'.join(clean(line) for line in text.splitlines())+'\n'
assert '{{' not in text
(OUT/'GuMaster2D_Project_Plan.md').write_text(text,encoding='utf-8')

fsd=['# GuMaster2D — Functional Specification Document','',
 'Phiên bản0.2,07/10/2026. Đề xuất cho closed alpha4người; Godot PC. Đây làđặc tảđểtriển khai, chưa phải phần mềmđãhoànthành.', '',
 'Cấu trúc thamkhảoDAUN-FSD[EN].xlsx; mởrộngAC/errors/input/output/trace. Rulebook,UC,data/API vàcatalog nằmtrongExcel.','']
for r in d.REQS:
    fsd += ['## '+r['id']+' — '+r['name'],'',
      '**Phân cấp:** '+r['module']+' → '+r['group']+' → '+r['screen'],
      '**Giai đoạn / ưu tiên / trạng thái:** '+r['phase']+' / '+r['priority']+' / Đề xuất','',
      r['spec'],'','**Input / tiền điều kiện:** '+r['inputs'],'',
      '**Output / hậu điều kiện:** '+r['outputs'],'',
      '**Validation / ngoại lệ:** '+r['errors'],'',
      '**Acceptance Criteria:** '+r['ac'],'',
      '**Trace:** WBS '+r['wbs']+'; Test '+r['test']+'; Rule '+(r['rule'] or 'Theo rulebook chung')+'; UC '+(r['uc'] or 'Luồng chuyên biệt trong FSD'), '']
fsd+=['# Non-functional Requirements','']
for n in d.NFR:fsd+=['## '+n[0]+' — '+n[1],'','**Điều kiện:** '+n[2],'','**Pass:** '+n[3],'','**Evidence:** '+n[5],'','**WBS:** '+n[4],'']
(OUT/'GuMaster2D_FSD.md').write_text('\n'.join(clean(x) for x in fsd),encoding='utf-8')
wbsmd=['# GuMaster2D — Work Breakdown Structure','',
 'Đề xuất0.2 /07/10/2026. 72workpackages;199PDbase;238.8PDgồmdựphòng. 1PD=8giờhữuích. Baseline36tuần/18sprint;20h/người/tuần;4người.','',
 'Dependency làfinish-to-start giữa cácsprint. Cùngsprintphốihợpbuild/integrate/test;gateđánhgiácuốisprint. LịchtạifileExcelbám00_START!B7. Test prototype báo cáo riêng;79testcase alpha chưa chạy. Ba gói spike S01 đang làm; code prototype tham khảo cho các gói sau không đồng nghĩa mở toàn bộ WIP hoặc hoàn thành AC alpha.','',
 md_table(['Code / ID','Sprint / tuần','Owner','Gói công việc','Predecessors','Base / Total PD','Status','Definition of Done'],[
 [t['code']+' / '+t['id'],f"S{t['sprint']:02} / {t['sw']}–{t['ew']}",t['role'],t['name'],t['deps'] or '—',f"{t['effort']:.1f} / {t['effort']*1.2:.1f}",'Đang làm (prototype)' if t['id'] in PROTOTYPE_TASKS else 'Chưa bắt đầu',t['dod']] for t in TASKS])]
(OUT/'GuMaster2D_WBS.md').write_text('\n'.join(clean(x) for x in wbsmd)+'\n',encoding='utf-8')

# Editable Word document generated from the same readable project plan.
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
doc=Document();sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Cm(1.7);sec.left_margin=sec.right_margin=Cm(1.7)
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
for name in ['Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor.from_string(NAVY)
    st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12)
doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 2'].font.size=Pt(14);doc.styles['Heading 3'].font.size=Pt(11)
header=sec.header.paragraphs[0];header.text='GUMASTER2D  /  PROJECT PLAN  •  FSD  •  WBS'
header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.text='Đề xuất v0.2  |  07/10/2026  |  Trang '
for r in footer.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(GRAY)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'vi-VN');normal.element.get_or_add_rPr().append(lang)

def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),BLUE);pr.append(color);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,line):
    # Keep real web links, bold spans and inline code readable in Word.
    pattern=r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for part in re.split(pattern,line):
        if not part:continue
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if m:hyperlink(p,m[1],m[2])
        elif part.startswith('**') and part.endswith('**'):p.add_run(part[2:-2]).bold=True
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
        else:p.add_run(part.replace('<br>','\n'))

lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i]
    if not line.strip():i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
        cells=[[x.strip() for x in row.strip('|').split('|')] for row in block]
        cells=[r for r in cells if not all(re.fullmatch(r'[:\- ]+',c or '-') for c in r)]
        table=doc.add_table(rows=1,cols=len(cells[0]));table.style='Table Grid';table.autofit=False
        available=17.6
        weights={2:[1,3],3:[1.1,3,6.5],4:[1.2,2.8,6.8,1],5:[1.8,3,3,3,3],6:[1,4,1,1,1,1]}.get(len(cells[0]),[1]*len(cells[0]))
        for j,w in enumerate(weights):table.columns[j].width=Cm(available*w/sum(weights))
        for ri,row in enumerate(cells):
            tr=table.rows[0] if ri==0 else table.add_row()
            no_split=OxmlElement('w:cantSplit');tr._tr.get_or_add_trPr().append(no_split)
            if ri==0:
                repeat=OxmlElement('w:tblHeader');tr._tr.get_or_add_trPr().append(repeat)
            for j,value in enumerate(row):
                cell=tr.cells[j];cell.width=Cm(available*weights[j]/sum(weights));p=cell.paragraphs[0];inline(p,value)
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.0
                for run in p.runs:
                    run.font.size=Pt(8.5)
                    if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),TEAL if ri==0 else ('EFF5F8' if ri%2 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shade)
        doc.add_paragraph();continue
    heading=re.match(r'^(#{1,4}) (.*)',line)
    if heading:
        level=len(heading[1]);p=doc.add_paragraph(style='Title' if level==1 else 'Heading '+str(level-1));inline(p,heading[2])
    else:p=doc.add_paragraph();inline(p,line)
    i+=1
doc.core_properties.title='GuMaster2D — Kế hoạch dự án, FSD và WBS'
doc.core_properties.subject='Đội 4 người tự phát triển; Godot PC Windows; closed alpha'
doc.core_properties.author='GuMaster2D planning'
DOCX=OUT/'GuMaster2D_Project_Plan.docx';doc.save(DOCX)
assert Document(DOCX).tables and len(Document(DOCX).paragraphs)>=90

# Implementation owns docs/game-studio; regenerating planning must preserve live reports.

readme='''# GuMaster2D — Bộ tài liệu lập kế hoạch

Ngày:07/10/2026. Đội4người tự pháttriển, chưa cóngânsách; Godot PC native trước; Web spike.

- `GuMaster2D_Project_Plan.docx`: bảnWord đểđọc vàtraođổi.
- `GuMaster2D_FSD_WBS_Plan.xlsx`:25sheet chi tiết, cócôngthức vàlịch.
- `GuMaster2D_Project_Plan.md`: bảnMarkdown củaWord.
- `GuMaster2D_FSD.md`: toàn bộ55yêucầuchứcnăng và12NFR.
- `GuMaster2D_WBS.md`:72workpackages,dependencies,effort vàDoD.
- `validation_report.json`: kếtquảkiểmtra tínhnhấtquántàiliệu; khôngphảitestgame.

Giả định:20h/người/tuần;75%hữuích;20%dựphòng;36tuần. Ngày12/10/2026–18/06/2027chỉminhhọa.
TrongExcel, sửa00_START!B7/B8/B9/B10 rồi xem19_RESOURCE và20_SPRINT_LOAD. Đổiavailability khôngtựnénWBS/Gantt; cầnrebaseline.

## Tạo lại tài liệu

Source:tools/planning_data.py vàtools/plan_template.md. Generator:tools/generate_plan.py.
Kế hoạch code Godot và trạng thái prototype: ../implementation/GODOT_IMPLEMENTATION_PLAN.md. CầnPython3.12+ vớiopenpyxl vàpython-docx. Generator khôngsửafileDAUNgốc.
Chạy: `python tools/generate_plan.py` từthưmụcnày.

Chạy generator sẽghiđècácartifactđãtạo. NếuđãchỉnhExcel/Word thủcông,hãylưubảncopyvàđưathayđổivềsource trướcchạylại.
Mọitestcase game ởtrạngtháiChưachạy; cácNFRchưađượcbenchmark. Cashsheetđơngiá0làinputchờđiền,khôngphảicamkếtmiễnphíhosting.
'''
(OUT/'README.md').write_text('\n'.join(clean(x) for x in readme.splitlines())+'\n',encoding='utf-8')
REPORT.update(word_paragraphs=len(Document(DOCX).paragraphs),word_tables=len(Document(DOCX).tables),
 artifacts=['GuMaster2D_FSD_WBS_Plan.xlsx','GuMaster2D_Project_Plan.docx','GuMaster2D_Project_Plan.md','GuMaster2D_FSD.md','GuMaster2D_WBS.md'])
(OUT/'validation_report.json').write_text(json.dumps(REPORT,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(REPORT,ensure_ascii=False,indent=2))
