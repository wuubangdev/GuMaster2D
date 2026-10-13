"""Build the editable Godot sprint workbook from the implementation plan.
Requires openpyxl. Regeneration overwrites generated XLSX/CSV, not the source plan.
"""
from pathlib import Path
import csv
import json
import math
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "implementation"
source = (OUT / "GODOT_IMPLEMENTATION_PLAN.md").read_text(encoding="utf-8")
tasks = []
reviewers = {"GD/QA": "CL", "CL": "SV", "SV": "CL", "AR": "GD/QA"}
for line in source.splitlines():
    if line.startswith("| G") and line[2:5].startswith("G0") or line.startswith("| G1"):
        row = [part.strip() for part in line.split("|")[1:-1]]
        if len(row) == 6 and row[0][1:].isdigit():
            tasks.append(dict(id=row[0], owner=row[1], task=row[2], hours=float(row[3]), deps=row[4], dod=row[5]))
assert len(tasks) == 11 and math.isclose(sum(t["hours"] for t in tasks), 84.0)

wb = Workbook()
wb.remove(wb.active)
wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
cache = {}

def sheet(name, title, headers, widths):
    ws = wb.create_sheet(name)
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).font = Font(size=18, bold=True, color="173D32")
    ws.row_dimensions[1].height = 32
    ws.append(["07/10/2026 | Godot4.7.2 / GDScript | PC native first | Giờ hữu ích, không phải giờ lịch"])
    ws.append([])
    ws.append(headers)
    for c in ws[4]:
        c.fill = PatternFill("solid", fgColor="214B3B")
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width
    ws.freeze_panes = "C5"
    ws.sheet_view.showGridLines = False
    return ws

def formula(ws, cell, expression, value):
    ws[cell] = expression
    cache[(ws.title, cell)] = value

ws = sheet("00_README", "GuMaster2D — Godot coding backlog", ["Mục", "Giá trị / cách dùng"], [30, 120])
notes = [
    ("Phạm vi", "11 task còn lại của sprint 2 tuần tiếp theo; 84h base +20%reserve =100,8h. Không cộng thêm vào tổng199PD alpha một cách máy móc."),
    ("Công suất", "4 người ×20h/tuần ×2tuần ×75%focus =120h hữu ích; giờ/role xem02_CAPACITY."),
    ("Status", "Tất cả task backlog mới chưa bắt đầu; prototype hiện có được trace riêng. Sửa Status/Evidence khi đội review."),
    ("FSD/WBS", "Đối chiếu ../planning/GuMaster2D_FSD_WBS_Plan.xlsx. Test local không nghiệm thu yêu cầu online."),
    ("Source", "GODOT_IMPLEMENTATION_PLAN.md; tools/generate_coding_backlog.py tạo lại XLSX/CSV. Lưu bản copy nếu đã sửa workbook thủ công."),
    ("Ngày", "Chưa chốt kickoff thật. Dùng ngày1–10 trong coding plan; lịch36tuần trong workbook chính vẫn chỉ là forecast."),
    ("Rủi ro", "Chưa re-estimate Godot/transport/DB; G04/G07 phối hợp bằng payload mẫu và stub ngày1, không chờ tuần tự toàn bộ."),
]
for row in notes:
    ws.append(row)

ws = sheet("01_BACKLOG", "Sprint Godot — 4 owners / 2 weeks", ["ID", "Owner", "Việc cụ thể", "Base h", "Reserve h", "Total h", "Dependency", "Reviewer", "Status", "Definition of Done", "Evidence / issue"], [12, 13, 62, 12, 13, 12, 30, 15, 20, 85, 50])
for rn, task in enumerate(tasks, 5):
    ws.append([task["id"], task["owner"], task["task"], task["hours"], None, None, task["deps"], reviewers[task["owner"]], "Chưa bắt đầu", task["dod"], ""])
    formula(ws, f"E{rn}", f"=D{rn}*0.2", task["hours"]*.2)
    formula(ws, f"F{rn}", f"=D{rn}+E{rn}", task["hours"]*1.2)
    ws.row_dimensions[rn].height = 68
ws.auto_filter.ref = "A4:K15"
status = DataValidation(type="list", formula1='"Chưa bắt đầu,Đang làm,Đang review,Hoàn thành,Bị chặn,Hoãn"')
ws.add_data_validation(status)
status.add("I5:I15")

ws = sheet("02_CAPACITY", "Công suất sprint — sửa giờ và focus khi đội chốt", ["Role", "h/tuần", "Focus", "Capacity h", "Base task h", "Incl reserve h", "Remaining h"], [16, 16, 14, 19, 19, 21, 20])
for rn, role in enumerate(["GD/QA", "CL", "SV", "AR"], 5):
    base = sum(t["hours"] for t in tasks if t["owner"] == role)
    ws.append([role, 20, .75, None, None, None, None])
    ws[f"C{rn}"].number_format = "0%"
    formula(ws, f"D{rn}", f"=B{rn}*2*C{rn}", 30)
    formula(ws, f"E{rn}", f'=SUMIF(\'01_BACKLOG\'!B5:B15,A{rn},\'01_BACKLOG\'!D5:D15)', base)
    formula(ws, f"F{rn}", f'=SUMIF(\'01_BACKLOG\'!B5:B15,A{rn},\'01_BACKLOG\'!F5:F15)', base*1.2)
    formula(ws, f"G{rn}", f"=D{rn}-F{rn}", 30-base*1.2)
ws.append(["TOTAL", "", "", 120, 84, 100.8, 19.2])
for col, value in [("D", 120), ("E", 84), ("F", 100.8), ("G", 19.2)]:
    formula(ws, f"{col}9", f"=SUM({col}5:{col}8)", value)
ws.append(["Lưu ý", "Thay capacity không tự đổi estimate/dependency; nếu role quá tải phải rebaseline."])

ws = sheet("03_FSD_TRACE", "Prototype evidence — tất cả chỉ một phần/local", ["FSD", "WBS", "Đã có", "Còn thiếu"], [30, 36, 88, 88])
part = source.split("## 4.")[1].split("## 5.")[0]
for line in part.splitlines():
    if line.startswith("| CL-") or line.startswith("| SV-"):
        ws.append([c.strip() for c in line.split("|")[1:-1]])

for ws in wb:
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.font = Font(name="Calibri", size=11, color="20382D")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c.row % 2:
                c.fill = PatternFill("solid", fgColor="EFF5EF")
        if ws.title != "01_BACKLOG":
            ws.row_dimensions[row[0].row].height = 58
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

path = OUT / "Godot_Coding_Backlog.xlsx"
wb.save(path)
# Cache formula values so file previews/data_only reads show the initial budget immediately.
ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", ns)
with zipfile.ZipFile(path) as source_zip:
    entries = {item.filename: source_zip.read(item.filename) for item in source_zip.infolist()}
for index, ws in enumerate(wb, 1):
    name = f"xl/worksheets/sheet{index}.xml"
    root = ET.fromstring(entries[name])
    for c in root.iter(f"{{{ns}}}c"):
        value = cache.get((ws.title, c.get("r")))
        if value is not None:
            v = c.find(f"{{{ns}}}v")
            if v is None:
                v = ET.SubElement(c, f"{{{ns}}}v")
            v.text = str(value)
    entries[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target_zip:
    for name, data in entries.items():
        target_zip.writestr(name, data)

with (OUT / "Godot_Coding_Backlog.csv").open("w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    writer.writerow(["ID", "Owner", "Task", "Base hours", "Hours with reserve", "Dependencies", "Status", "DoD"])
    for t in tasks:
        writer.writerow([t["id"], t["owner"], t["task"], t["hours"], t["hours"]*1.2, t["deps"], "Chưa bắt đầu", t["dod"]])
checked = load_workbook(path, data_only=True)
assert checked["01_BACKLOG"].max_row == 15
assert math.isclose(sum(checked["02_CAPACITY"].cell(r,6).value for r in range(5,9)), 100.8)
print(json.dumps({"tasks":len(tasks), "base_hours":84, "with_reserve_hours":100.8, "capacity_hours":120, "status":"PASS"}))
